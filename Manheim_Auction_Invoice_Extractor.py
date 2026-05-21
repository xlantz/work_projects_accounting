#!/usr/bin/env python3
"""
Auction Proceed Invoice Extractor — Manheim

Column definitions:
  Gross  = vehicle sale amount (absolute value)
  Net    = Total Payments received
  Auct   = Gross - Net  (Excel formula)
  Sell   = sum of sell-fee line items
  Recon  = Auct - Sell  (Excel formula)
"""

import os
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import pdfplumber
except ImportError:
    print("Installing pdfplumber...")
    os.system(f"{sys.executable} -m pip install pdfplumber --break-system-packages -q")
    import pdfplumber

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl --break-system-packages -q")
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Sell-fee keywords — ONLY these count as sell fees (case-insensitive)
# ---------------------------------------------------------------------------
SELL_FEE_KEYWORDS = [
    "sell fee",
    "seller success fee",
    "simulcast seller premium fee",
]


def is_sell_fee(description: str) -> bool:
    return any(kw in description.lower().strip() for kw in SELL_FEE_KEYWORDS)


def parse_amount(raw: str) -> float:
    raw = raw.strip().replace(",", "").replace("$", "").replace(" ", "")
    negative = raw.startswith("(") or raw.startswith("-")
    raw = raw.strip("()").lstrip("-")
    try:
        val = float(raw)
    except ValueError:
        val = 0.0
    return -val if negative else val


def parse_date(raw: str):
    for fmt in ("%d-%b-%Y", "%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return datetime.strptime(raw.strip(), fmt).date()
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------------------
# Text extraction — handles both real PDFs and Manheim zip-wrapped PDFs
# ---------------------------------------------------------------------------

def is_zip_file(path: str) -> bool:
    try:
        with open(path, "rb") as f:
            return f.read(4) == b"PK\x03\x04"
    except Exception:
        return False


def extract_text(pdf_path: str) -> str:
    """Extract text from a Manheim invoice — either a real PDF or a zip-wrapped one."""
    if is_zip_file(pdf_path):
        # Zip format: contains a 1.txt with the invoice text
        with zipfile.ZipFile(pdf_path, "r") as z:
            all_txt = [n for n in z.namelist() if n.endswith(".txt") and n != "manifest.json"]
            if not all_txt:
                raise ValueError(f"No .txt found inside zip: {pdf_path}")
            page1 = next((n for n in all_txt if re.search(r"\b1\.txt$", n)), sorted(all_txt)[0])
            with z.open(page1) as f:
                return f.read().decode("utf-8", errors="replace")
    else:
        # Real PDF — use pdfplumber
        text = ""
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text


# ---------------------------------------------------------------------------
# Manheim parser
# ---------------------------------------------------------------------------

def parse_invoice(text: str) -> dict:
    data = {
        "credit_memo": "",
        "loan_number": "",
        "auction_house": "",
        "invoice_date": None,
        "gross": 0.0,
        "net": 0.0,
        "sell_fees": 0.0,
    }

    # Credit Memo
    m = re.search(r"Credit Memo #\s*(\d+)", text, re.IGNORECASE)
    if m:
        data["credit_memo"] = m.group(1)

    # Lease Account No — fallback to VIN
    m = re.search(r"LEASE ACCOUNT NO\s+(\d+)", text, re.IGNORECASE)
    if m:
        data["loan_number"] = m.group(1)
    if not data["loan_number"]:
        m = re.search(r"VIN\s+([A-HJ-NPR-Z0-9]{17})", text, re.IGNORECASE)
        if m:
            data["loan_number"] = m.group(1)

    # Invoice Date
    m = re.search(r"INVOICE DATE\s+(\S+)", text, re.IGNORECASE)
    if m:
        data["invoice_date"] = parse_date(m.group(1))

    # Auction house: strip leading station number, handle multi-word name wrapping
    # e.g. "541 MANHEIM DENVER", or "MANHEIM NORTH\nCAROLINA", "MANHEIM BALTIMORE-\nWASHINGTON"
    loc_m = re.search(
        r"TRANSACTION LOCATION\s*[\r\n]+(.*?)[\r\n]+(.*?)[\r\n]",
        text, re.IGNORECASE,
    )
    if not loc_m:
        loc_m = re.search(r"TRANSACTION LOCATION\s+(.+)", text, re.IGNORECASE)
        if loc_m:
            data["auction_house"] = re.sub(r"^\d+\s+", "", loc_m.group(1).strip())
    else:
        line1 = re.sub(r"^\d+\s+", "", loc_m.group(1).strip())
        line2 = loc_m.group(2).strip()
        is_continuation = (
            line2
            and not re.match(r"\d{2}-[A-Z]{3}-\d{4}", line2)
            and not re.match(r"\d+\s+[A-Z]", line2)
            and not re.search(r"(INVOICING|LOCATION|ACCOUNT|BUYER|LEASE|MILEAGE|VIN|YEAR)", line2, re.IGNORECASE)
            and not re.match(r"\d{4,}", line2)
            and re.match(r"^[A-Z][A-Z\s\-]+$", line2)
        )
        if is_continuation:
            line1 = (line1.rstrip("-").rstrip() + " " + line2).strip()
        data["auction_house"] = line1

    # Charge table only (between "PRICE TAX LINE TOTAL" and "ADJUSTMENTS")
    charge_match = re.search(
        r"PRICE\s+TAX\s+LINE TOTAL\s*\n(.*?)(?=\nADJUSTMENTS|\nPAYMENTS)",
        text, re.DOTALL | re.IGNORECASE,
    )
    charge_lines = []
    if charge_match:
        charge_lines = [
            l.strip()
            for l in charge_match.group(1).replace("\r\n", "\n").replace("\r", "\n").split("\n")
            if l.strip()
        ]

    amount_pat = re.compile(r"\(?\$[\d,]+\.\d{2}\)?$")
    skip = {"TOTAL", "SUB TOTAL", "TAX", "ADJUSTMENTS", "PAYMENTS", "AMOUNT DUE"}

    for line in charge_lines:
        m = amount_pat.search(line)
        if not m:
            continue
        amount_str = m.group(0)
        desc_area = line[: m.start()].strip()
        if any(p in desc_area.upper() for p in skip):
            continue
        amount = parse_amount(amount_str)
        uk_m = re.search(r"\d{4}-\d+-\d+-\d+\s+(.*)", desc_area)
        description = uk_m.group(1).strip() if uk_m else desc_area
        description = re.sub(r"\s+\$[\d,]+\.?\d*$", "", description).strip()

        if amount < 0:
            data["gross"] = abs(amount)
        elif is_sell_fee(description):
            data["sell_fees"] += amount

    # Net = Total Payments
    m = re.search(r"TOTAL PAYMENTS\s+\(?\$?([\d,]+\.\d{2})\)?", text, re.IGNORECASE)
    if m:
        data["net"] = parse_amount(m.group(1))

    return data


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

HEADERS = ["Loan #", "Gross", "Net", "Auct", "Sell", "Recon", "Credit Memo", "Auction House", "Date", "Note for Account"]
COL_WIDTHS = [12, 10, 10, 10, 10, 10, 14, 22, 12, 20]


def apply_header_style(cell):
    cell.font = Font(bold=True, name="Arial", size=10)
    cell.fill = PatternFill("solid", start_color="D9E1F2")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def apply_data_style(cell, is_currency=False, is_date=False):
    thin = Side(style="thin")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.font = Font(name="Arial", size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if is_currency:
        cell.number_format = '$#,##0.00'
    if is_date:
        cell.number_format = 'MM/DD/YYYY'


def write_excel(records: list, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    tab_name = datetime.today().strftime("%m.%d.%y")
    ws.title = tab_name

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        apply_header_style(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS[col_idx - 1]
    ws.row_dimensions[1].height = 20

    for row_idx, rec in enumerate(records, start=2):
        r = row_idx
        values = [
            rec["loan_number"],
            rec["gross"],
            rec["net"],
            f"=B{r}-C{r}",   # Auct = Gross - Net
            rec["sell_fees"],
            f"=D{r}-E{r}",   # Recon = Auct - Sell
            rec["credit_memo"],
            rec["auction_house"],
            rec["invoice_date"],
            "",
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=val)
            apply_data_style(cell, is_currency=col_idx in {2, 3, 4, 5, 6}, is_date=col_idx == 9)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"\n✅ Excel saved: {output_path}  (tab: '{tab_name}')")


# ---------------------------------------------------------------------------
# PDF renaming
# ---------------------------------------------------------------------------

def rename_pdfs(folder: str, records: list, pdf_files: list):
    print("\nRenaming PDF files...")
    for pdf_file, rec in zip(pdf_files, records):
        loan = rec.get("loan_number", "").strip()
        if not loan:
            print(f"  ⚠️  Skipping rename for {pdf_file} — no Loan # found")
            continue
        old_path = os.path.join(folder, pdf_file)
        new_name = f"{loan}Auct.pdf"
        new_path = os.path.join(folder, new_name)
        if os.path.abspath(old_path) == os.path.abspath(new_path):
            print(f"  ✓  {pdf_file} already named correctly")
            continue
        if os.path.exists(new_path):
            print(f"  ⚠️  {new_name} already exists — skipping rename for {pdf_file}")
            continue
        os.rename(old_path, new_path)
        print(f"  ✓  {pdf_file}  →  {new_name}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print("=" * 60)
    print("  Auction Proceed Invoice Extractor — Manheim")
    print("=" * 60)

    folder = input("\nEnter the folder path containing the PDF invoice files: ").strip().strip('"').strip("'")
    if not folder:
        print("No folder entered. Exiting.")
        sys.exit(1)
    if not os.path.isdir(folder):
        print(f"Error: '{folder}' is not a valid directory.")
        sys.exit(1)

    pdf_files = sorted([f for f in os.listdir(folder) if f.lower().endswith(".pdf")])
    if not pdf_files:
        print(f"No PDF files found in '{folder}'.")
        sys.exit(1)

    print(f"\nFound {len(pdf_files)} PDF file(s):")
    for f in pdf_files:
        print(f"  • {f}")

    records = []
    failed = []

    for pdf_file in pdf_files:
        pdf_path = os.path.join(folder, pdf_file)
        print(f"\nProcessing: {pdf_file} ...", end=" ")
        try:
            text = extract_text(pdf_path)
            rec = parse_invoice(text)
            records.append(rec)
            print(
                f"✓  Loan #{rec['loan_number']}  |  "
                f"Gross ${rec['gross']:,.2f}  |  Net ${rec['net']:,.2f}  |  "
                f"Sell ${rec['sell_fees']:,.2f}"
            )
        except Exception as e:
            print(f"✗  ERROR: {e}")
            failed.append(pdf_file)
            records.append({
                "credit_memo": "", "loan_number": "", "auction_house": "",
                "invoice_date": None, "gross": 0.0, "net": 0.0, "sell_fees": 0.0,
            })

    if failed:
        print(f"\n⚠️  Failed to parse {len(failed)} file(s): {', '.join(failed)}")

    today_str = datetime.today().strftime("%m.%d.%y")
    output_excel = os.path.join(folder, f"AuctionProceeds_{today_str}.xlsx")
    write_excel(records, output_excel)
    rename_pdfs(folder, records, pdf_files)

    print("\nDone! 🎉")


if __name__ == "__main__":
    main()